import io
import os
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook


MAX_EXTRACTED_CHARS = 40000
MAX_DOCUMENT_FILE_SIZE = 5 * 1024 * 1024
SUPPORTED_EXTENSIONS = (".docx", ".xlsx")
TRUNCATION_NOTICE = "[The document content was truncated]"

WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class DocumentReadError(RuntimeError):
    pass


def is_supported_document(filename):
    if not isinstance(filename, str):
        return False
    return os.path.splitext(filename)[1].lower() in SUPPORTED_EXTENSIONS


def extract_document_text(filename, data):
    extension = os.path.splitext(filename if isinstance(filename, str) else "")[1].lower()
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise DocumentReadError("The document is empty.")
    if extension == ".xlsx":
        text = _extract_xlsx(data)
    elif extension == ".docx":
        text = _extract_docx(data)
    else:
        raise DocumentReadError(f"Unsupported document type: {extension or 'unknown'}")
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS] + "\n" + TRUNCATION_NOTICE
    return text


def build_document_user_text(filename, extracted_text, caption=None):
    parts = [
        f'The customer sent a document "{filename}". '
        "The system converted the document to plain text below.",
        "--- Converted document content ---",
        extracted_text,
        "--- End of document content ---",
    ]
    if isinstance(caption, str) and caption.strip():
        parts.append(f"Customer's message with the document: {caption.strip()}")
    return "\n".join(parts)


def _extract_xlsx(data):
    try:
        workbook = load_workbook(
            io.BytesIO(bytes(data)),
            read_only=True,
            data_only=True,
        )
    except Exception:
        raise DocumentReadError("The Excel file could not be parsed.")
    try:
        sections = []
        for sheet in workbook.worksheets:
            rows = _read_xlsx_rows(sheet)
            if rows:
                sections.append(f"--- Sheet: {sheet.title} ---\n" + "\n".join(rows))
    finally:
        workbook.close()
    if not sections:
        raise DocumentReadError("The Excel file has no text content.")
    return "\n\n".join(sections)


def _read_xlsx_rows(sheet):
    rows = []
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        while values and _is_empty_cell(values[-1]):
            values.pop()
        if values and not all(_is_empty_cell(value) for value in values):
            rows.append("\t".join(_cell_text(value) for value in values))
    return rows


def _is_empty_cell(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _cell_text(value):
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _extract_docx(data):
    try:
        with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
            document_xml = archive.read("word/document.xml")
        root = ET.fromstring(document_xml)
    except (zipfile.BadZipFile, KeyError, ET.ParseError):
        raise DocumentReadError("The Word document could not be parsed.")
    body = root.find(_word_tag("body"))
    if body is None:
        raise DocumentReadError("The Word document could not be parsed.")
    lines = []
    _read_word_body(body, lines)
    content = "\n".join(lines).strip()
    if not content:
        raise DocumentReadError("The Word document has no text content.")
    return content


def _read_word_body(element, lines):
    for child in element:
        if child.tag == _word_tag("p"):
            text = _read_word_paragraph(child)
            if text.strip():
                lines.append(text)
        elif child.tag == _word_tag("tbl"):
            rows = _read_word_table(child)
            lines.extend(row for row in rows if row.strip())
        else:
            _read_word_body(child, lines)


def _read_word_paragraph(paragraph):
    chunks = []
    for node in paragraph.iter():
        if node.tag == _word_tag("t"):
            chunks.append(node.text or "")
        elif node.tag == _word_tag("tab"):
            chunks.append("\t")
        elif node.tag == _word_tag("br"):
            chunks.append("\n")
    return "".join(chunks)


def _read_word_table(table):
    rows = []
    for row in table.findall(_word_tag("tr")):
        cells = []
        for cell in row.findall(_word_tag("tc")):
            texts = []
            for paragraph in cell.iter(_word_tag("p")):
                text = _read_word_paragraph(paragraph).replace("\t", " ").strip()
                if text:
                    texts.append(text)
            cells.append(" ".join(texts))
        rows.append("\t".join(cells))
    return rows


def _word_tag(name):
    return f"{{{WORD_NAMESPACE}}}{name}"
